import io
import zipfile
from decimal import Decimal

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter
from pypdf.generic import (
    ArrayObject,
    DecodedStreamObject,
    DictionaryObject,
    NameObject,
    TextStringObject,
)

from pc_manager_agent.config.office import OfficeLimits
from pc_manager_agent.domain.office_documents import (
    DocumentFormat,
    DocumentSupport,
    OfficeError,
    OfficeValue,
    ValueKind,
)
from pc_manager_agent.office.adapters import parse_document
from pc_manager_agent.office.text import json_set, strict_json
from pc_manager_agent.safety.office.content import inspect_package, safe_spreadsheet_text


@pytest.mark.parametrize("content", ["<definedName/>", "<r><t>rich</t></r>", "<extLst/>"])
def test_unmodeled_spreadsheet_features_are_read_only(content):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("_rels/.rels", "<Relationships/>")
        archive.writestr("xl/synthetic.XML", f"<root>{content}</root>")
    assert "LIMITED_EDIT_SUPPORT" in inspect_package(
        buffer.getvalue(), DocumentFormat.XLSX, OfficeLimits()
    )


def pdf(*, text=True, encrypted=False, active=False):
    writer = PdfWriter()
    page = writer.add_blank_page(300, 300)
    if text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        page[NameObject("/Resources")] = DictionaryObject(
            {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
        )
        content = DecodedStreamObject()
        content.set_data(b"BT /F1 12 Tf 10 100 Td (Synthetic PDF text) Tj ET")
        page[NameObject("/Contents")] = content
    if active:
        writer._root_object[NameObject("/OpenAction")] = ArrayObject(
            [TextStringObject("NEVER_EXECUTE")]
        )
    if encrypted:
        writer.encrypt("synthetic-password")
    stream = io.BytesIO()
    writer.write(stream)
    writer.close()
    return stream.getvalue()


def test_pdf_static_text_blank_ocr_encryption_and_limits():
    result = parse_document(pdf(active=True), DocumentFormat.PDF, OfficeLimits())
    assert result.blocks[0].text == "Synthetic PDF text"
    assert result.blocks[0].page == 1
    assert result.support is DocumentSupport.READ_ONLY
    assert "PDF_ACTIVE_CONTENT_IGNORED" in result.warnings
    blank = parse_document(pdf(text=False), DocumentFormat.PDF, OfficeLimits())
    assert "OCR_REQUIRED_OR_EMPTY_PAGE" in blank.warnings
    for data, limits in (
        (pdf(encrypted=True), OfficeLimits()),
        (pdf(), OfficeLimits(max_text_chars=1)),
        (b"%PDF-malformed", OfficeLimits()),
        (b"wrong header", OfficeLimits()),
    ):
        with pytest.raises(OfficeError):
            parse_document(data, DocumentFormat.PDF, limits)


def test_json_edit_never_rounds_unedited_decimal_values():
    result = json_set(
        '{"precise":1.234567890123456789,"list":[1.01, true, null],"name":"before"}',
        "/name",
        OfficeValue(value="after"),
    )
    parsed = strict_json(result)
    assert parsed["precise"] == Decimal("1.234567890123456789")
    assert parsed["list"] == [Decimal("1.01"), True, None]
    assert parsed["name"] == "after"
    result = json_set(
        '{"value":1.1}', "/value", OfficeValue(kind=ValueKind.NUMBER, value="2.1234567890123456789")
    )
    assert strict_json(result)["value"] == Decimal("2.1234567890123456789")
    with pytest.raises(OfficeError):
        json_set('{"a":[{"x":1}]}', "/a/-1/x", OfficeValue(value="no"))


@pytest.mark.parametrize("value", ["\u200b=HYPERLINK(1)", "\u2028=1+1", "\ufeff\u200b +1"])
def test_unicode_formula_prefixes_remain_text(value):
    assert safe_spreadsheet_text(value).startswith("'")


def test_docx_multirun_and_xlsx_protection_are_not_silently_saved():
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("first")
    paragraph.add_run("second").bold = True
    stream = io.BytesIO()
    document.save(stream)
    assert (
        parse_document(stream.getvalue(), DocumentFormat.DOCX, OfficeLimits()).support
        is DocumentSupport.READ_ONLY
    )
    book = Workbook()
    book.active["A1"] = "ordinary"
    book.active.protection.sheet = True
    stream = io.BytesIO()
    book.save(stream)
    book.close()
    assert (
        parse_document(stream.getvalue(), DocumentFormat.XLSX, OfficeLimits()).support
        is DocumentSupport.READ_ONLY
    )
