from __future__ import annotations

import io
import zipfile

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from pc_manager_agent.config.office import OfficeLimits
from pc_manager_agent.domain.office_documents import (
    DocumentFormat,
    DocumentSupport,
    OfficeError,
    OfficeValue,
    ValueKind,
)
from pc_manager_agent.office.adapters import parse_document, serialize_document
from pc_manager_agent.office.text import json_set, strict_json
from pc_manager_agent.safety.office.content import (
    require_safe_formula,
    safe_spreadsheet_text,
    safe_xml,
)


@pytest.mark.parametrize(
    "format_",
    [DocumentFormat.TXT, DocumentFormat.MARKDOWN, DocumentFormat.JSON, DocumentFormat.CSV],
)
def test_text_roundtrip(format_: DocumentFormat) -> None:
    data = b'{"name":"text"}' if format_ is DocumentFormat.JSON else "标题,值\r\n甲,1\r\n".encode()
    document = parse_document(data, format_, OfficeLimits())
    result = parse_document(serialize_document(document), format_, OfficeLimits())
    assert result == document


def test_simple_docx_roundtrip() -> None:
    source = Document()
    source.add_heading("报告", 1)
    source.add_paragraph("Unicode 内容")
    source.add_table(rows=1, cols=2).cell(0, 0).text = "表格"
    stream = io.BytesIO()
    source.save(stream)
    data = stream.getvalue()
    parsed = parse_document(data, DocumentFormat.DOCX, OfficeLimits())
    assert parsed.support is DocumentSupport.EDITABLE
    assert parsed.blocks[0].heading_level == 1
    assert (
        parse_document(serialize_document(parsed, data), DocumentFormat.DOCX, OfficeLimits())
        == parsed
    )


def test_xlsx_types_and_formula_preservation() -> None:
    source = Workbook()
    sheet = source.active
    assert sheet is not None
    sheet["A1"] = "name"
    sheet["A2"] = 12.5
    sheet["B2"] = "=SUM(A2:A2)"
    sheet["C2"] = True
    stream = io.BytesIO()
    source.save(stream)
    source.close()
    data = stream.getvalue()
    parsed = parse_document(data, DocumentFormat.XLSX, OfficeLimits())
    assert parsed.support is DocumentSupport.EDITABLE
    assert parsed.formula_count == 1
    assert (
        parse_document(serialize_document(parsed, data), DocumentFormat.XLSX, OfficeLimits())
        == parsed
    )


def test_xlsx_plain_text_equals_is_not_formula() -> None:
    source = Workbook()
    assert source.active is not None
    source.active["A1"] = "text"
    stream = io.BytesIO()
    source.save(stream)
    source.close()
    parsed = parse_document(stream.getvalue(), DocumentFormat.XLSX, OfficeLimits())
    cell = (
        parsed.sheets[0]
        .cells[0]
        .model_copy(update={"value": OfficeValue(value="=HYPERLINK(test)")})
    )
    sheet = parsed.sheets[0].model_copy(update={"cells": (cell,)})
    parsed = parsed.model_copy(update={"sheets": (sheet,)})
    result = load_workbook(io.BytesIO(serialize_document(parsed)), data_only=False)
    assert result.active is not None
    assert result.active["A1"].data_type == "s"
    result.close()


@pytest.mark.parametrize(
    "value", ["=cmd", "+SUM(1)", "-HYPERLINK(a)", "@test", " \t=1+1", "\rtest", "＝1+1"]
)
def test_csv_formula_text_escape(value: str) -> None:
    assert safe_spreadsheet_text(value) == "'" + value


@pytest.mark.parametrize(
    "value", ["=WEBSERVICE(A1)", "=HYPERLINK(A1)", "=[x]A1", "=A1!B1", "=cmd|' /C calc'!A0"]
)
def test_unsafe_formula_rejected(value: str) -> None:
    with pytest.raises(OfficeError):
        require_safe_formula(value)


def test_json_is_strict_and_pointer_is_not_code() -> None:
    for text in ('{"a":1,"a":2}', '{"a":NaN}', '{"a":Infinity}', '{"a":1e1000}'):
        with pytest.raises(OfficeError):
            strict_json(text)
    assert strict_json(
        json_set(
            '{"a":1}',
            "/a",
            OfficeValue(
                kind=ValueKind.NUMBER,
                value="2",
            ),
        )
    ) == {"a": 2}
    with pytest.raises(OfficeError):
        json_set('{"a":1}', "/missing", OfficeValue(value="x"))


def test_entities_are_rejected() -> None:
    with pytest.raises(OfficeError):
        safe_xml(b'<!DOCTYPE x [<!ENTITY y "evil">]><x>&y;</x>')


def test_zip_traversal_is_rejected() -> None:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("../escape.xml", "<x/>")
    with pytest.raises(OfficeError, match="UNSAFE_PACKAGE_ENTRY"):
        parse_document(stream.getvalue(), DocumentFormat.DOCX, OfficeLimits())


def test_size_budget_and_legacy_document() -> None:
    with pytest.raises(OfficeError, match="SIZE_LIMIT"):
        parse_document(b"long", DocumentFormat.TXT, OfficeLimits(text_bytes=2))
    with pytest.raises(OfficeError, match="LEGACY"):
        parse_document(b"\xd0\xcf\x11\xe0data", DocumentFormat.DOCX, OfficeLimits())


@pytest.mark.parametrize("value", ["NaN", "Infinity", "1e1000", "not-a-number"])
def test_numbers_must_be_finite_and_bounded(value: str) -> None:
    with pytest.raises(ValueError):
        OfficeValue(kind=ValueKind.NUMBER, value=value)
