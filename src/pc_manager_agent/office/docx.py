"""Conservative DOCX adapter; no Office process, macros, COM or external resource access."""

from __future__ import annotations

import io
import zipfile

from docx import Document

from pc_manager_agent.config.office import OfficeLimits
from pc_manager_agent.domain.office_documents import (
    DocumentBlock,
    DocumentCell,
    DocumentFormat,
    DocumentSheet,
    DocumentSupport,
    OfficeError,
    OfficeValue,
    StructuredDocument,
    ValueKind,
)
from pc_manager_agent.safety.office.content import safe_xml


def parse_docx(
    data: bytes, format_: DocumentFormat, warnings: tuple[str, ...], limits: OfficeLimits
) -> StructuredDocument:
    """Extract simple paragraphs/tables and explicitly restrict complex structures."""
    if format_ is DocumentFormat.DOCM:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = safe_xml(archive.read("word/document.xml"))
        texts = tuple(
            DocumentBlock(
                reference=f"p:{index}",
                text="".join(child.text or "" for child in node.iter() if child.tag.endswith("}t")),
            )
            for index, node in enumerate(node for node in root.iter() if node.tag.endswith("}p"))
        )
        return StructuredDocument(
            format=format_, blocks=texts, support=DocumentSupport.READ_ONLY, warnings=warnings
        )
    document = Document(io.BytesIO(data))
    flags = set(warnings)
    blocks: list[DocumentBlock] = []
    sheets: list[DocumentSheet] = []
    count = 0
    for index, paragraph in enumerate(document.paragraphs):
        style_name = paragraph.style.name if paragraph.style is not None else ""
        heading = (
            int(style_name[-1])
            if style_name.startswith("Heading ") and style_name[-1:].isdigit()
            else 0
        )
        if len(paragraph.runs) > 1:
            flags.add("LIMITED_EDIT_SUPPORT")
        blocks.append(
            DocumentBlock(reference=f"p:{index}", text=paragraph.text, heading_level=heading)
        )
        count += len(paragraph.text)
    for index, table in enumerate(document.tables):
        cells: list[DocumentCell] = []
        for row_index, row in enumerate(table.rows, 1):
            for column, cell in enumerate(row.cells, 1):
                if len(cell.paragraphs) > 1 or cell.tables:
                    flags.add("LIMITED_EDIT_SUPPORT")
                count += len(cell.text)
                cells.append(
                    DocumentCell(
                        reference=f"t:{index}:r:{row_index}:c:{column}",
                        row=row_index,
                        column=column,
                        value=OfficeValue(value=cell.text),
                    )
                )
        if len(cells) > limits.max_cells:
            raise OfficeError("TABLE_LIMIT")
        sheets.append(
            DocumentSheet(
                reference=f"t:{index}",
                name=f"Table {index + 1}",
                cells=tuple(cells),
                rows=len(table.rows),
                columns=len(table.columns),
            )
        )
    if count > limits.max_text_chars or len(sheets) > limits.max_sheets:
        raise OfficeError("DOCUMENT_LIMIT")
    return StructuredDocument(
        format=format_,
        blocks=tuple(blocks),
        sheets=tuple(sheets),
        warnings=tuple(sorted(flags)),
        support=DocumentSupport.READ_ONLY if flags else DocumentSupport.EDITABLE,
    )


def serialize_docx(document: StructuredDocument, original: bytes | None) -> bytes:
    """Apply only explicit supported paragraph/table values; keep untouched simple formatting."""
    output = Document(io.BytesIO(original)) if original is not None else Document()
    existing = list(output.paragraphs)
    if len(document.blocks) < len(existing):
        raise OfficeError("PARAGRAPH_REMOVAL_UNSUPPORTED")
    for index, block in enumerate(document.blocks):
        if index < len(existing):
            paragraph = existing[index]
            if paragraph.text != block.text:
                if len(paragraph.runs) > 1:
                    raise OfficeError("COMPLEX_PARAGRAPH_EDIT_BLOCKED")
                if paragraph.runs:
                    paragraph.runs[0].text = block.text
                else:
                    paragraph.add_run(block.text)
        else:
            paragraph = output.add_paragraph(block.text)
        if block.heading_level:
            paragraph.style = f"Heading {block.heading_level}"
    if len(document.sheets) < len(output.tables):
        raise OfficeError("TABLE_REMOVAL_UNSUPPORTED")
    for index, sheet in enumerate(document.sheets):
        table = (
            output.tables[index]
            if index < len(output.tables)
            else output.add_table(rows=sheet.rows, cols=sheet.columns)
        )
        if len(table.rows) != sheet.rows or len(table.columns) != sheet.columns:
            raise OfficeError("TABLE_STRUCTURE_CHANGE_UNSUPPORTED")
        for cell in sheet.cells:
            if cell.value.kind is not ValueKind.TEXT:
                raise OfficeError("DOCX_TABLE_REQUIRES_TEXT")
            target = table.cell(cell.row - 1, cell.column - 1)
            if target.text != cell.value.value:
                target.text = cell.value.value
    stream = io.BytesIO()
    output.save(stream)
    return stream.getvalue()
