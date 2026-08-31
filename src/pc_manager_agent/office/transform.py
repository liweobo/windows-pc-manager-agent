"""Finite deterministic document transformations and source-addressed differences."""

from __future__ import annotations

from decimal import ROUND_HALF_EVEN, Decimal, localcontext

from openpyxl.utils.cell import get_column_letter, range_boundaries

from pc_manager_agent.config.office import OfficeLimits
from pc_manager_agent.domain.office_documents import (
    DocumentBlock,
    DocumentCell,
    DocumentFormat,
    DocumentSheet,
    DocumentSupport,
    OfficeError,
    OfficeValue,
    StructuredDocument,
    ValueKind,
)
from pc_manager_agent.domain.office_plans import (
    DocumentDifference,
    DocumentEditPlan,
    DocumentOperation,
    DocumentOperationKind,
    OutputMode,
)
from pc_manager_agent.office.text import json_set
from pc_manager_agent.safety.office.content import require_safe_formula, safe_spreadsheet_text


def transform_documents(
    sources: tuple[StructuredDocument, ...],
    plan: DocumentEditPlan,
    limits: OfficeLimits,
) -> StructuredDocument:
    """Compile explicit edits; conversion never silently carries active Office structures."""
    if len(plan.operations) > limits.max_operations:
        raise OfficeError("OPERATION_LIMIT")
    if plan.output.mode is OutputMode.CREATE_NEW:
        if plan.initial_document is None:
            raise OfficeError("INITIAL_DOCUMENT_REQUIRED")
        document = plan.initial_document
    else:
        if not sources:
            raise OfficeError("DOCUMENT_SOURCE_REQUIRED")
        document = sources[0]
        if len(sources) > 1:
            document = merge_documents(sources, plan.output.format, limits)
        elif document.format is not plan.output.format:
            document = convert_document(document, plan.output.format)
        elif document.support is not DocumentSupport.EDITABLE:
            raise OfficeError("DOCUMENT_READ_ONLY")
    if (
        document.format is not plan.output.format
        or document.support is not DocumentSupport.EDITABLE
    ):
        raise OfficeError("OUTPUT_FORMAT_MISMATCH")
    targets: set[str] = set()
    for operation in plan.operations:
        if operation.target in targets:
            raise OfficeError("OVERLAPPING_EDIT_TARGET")
        targets.add(operation.target)
        document = apply_operation(document, operation)
    if document.format is DocumentFormat.CSV:
        sheet = document.sheets[0]
        safe_cells = tuple(
            cell.model_copy(
                update={"value": OfficeValue(value=safe_spreadsheet_text(cell.value.value))}
            )
            for cell in sheet.cells
        )
        document = document.model_copy(
            update={"sheets": (sheet.model_copy(update={"cells": safe_cells}),)}
        )
    if sum(len(sheet.cells) for sheet in document.sheets) > limits.max_cells:
        raise OfficeError("CELL_LIMIT")
    if sum(len(block.text) for block in document.blocks) > limits.max_text_chars:
        raise OfficeError("TEXT_LIMIT")
    if len(document.sheets) > limits.max_sheets:
        raise OfficeError("SHEET_LIMIT")
    document = document.model_copy(
        update={
            "formula_count": sum(
                cell.value.kind is ValueKind.FORMULA
                for sheet in document.sheets
                for cell in sheet.cells
            )
        }
    )
    return StructuredDocument.model_validate_json(document.model_dump_json())


def convert_document(source: StructuredDocument, target: DocumentFormat) -> StructuredDocument:
    """Create only documented derived formats; preserve sources as explicit text references."""
    if target is DocumentFormat.DOCX and source.format in {
        DocumentFormat.TXT,
        DocumentFormat.MARKDOWN,
        DocumentFormat.PDF,
    }:
        blocks: list[DocumentBlock] = []
        for block in source.blocks:
            if block.page is not None:
                blocks.append(
                    DocumentBlock(
                        reference=f"p:{len(blocks)}",
                        text=f"[Source page {block.page}]",
                    )
                )
            for line in block.text.splitlines():
                level = len(line) - len(line.lstrip("#"))
                heading = (
                    level if source.format is DocumentFormat.MARKDOWN and 1 <= level <= 6 else 0
                )
                blocks.append(
                    DocumentBlock(
                        reference=f"p:{len(blocks)}",
                        text=line[heading:].lstrip() if heading else line,
                        heading_level=heading,
                    )
                )
        return StructuredDocument(format=target, blocks=tuple(blocks))
    if target is DocumentFormat.XLSX and source.format is DocumentFormat.CSV:
        sheet = source.sheets[0]
        cells = tuple(
            cell.model_copy(update={"reference": f"s:0:{get_column_letter(cell.column)}{cell.row}"})
            for cell in sheet.cells
        )
        return StructuredDocument(
            format=target, sheets=(sheet.model_copy(update={"name": "Data", "cells": cells}),)
        )
    if (
        target in {DocumentFormat.TXT, DocumentFormat.MARKDOWN}
        and source.format is DocumentFormat.PDF
    ):
        return StructuredDocument(
            format=target,
            blocks=(
                DocumentBlock(
                    reference="body",
                    text="\n\n".join(
                        f"[Source page {block.page}]\n{block.text}" for block in source.blocks
                    ),
                ),
            ),
        )
    raise OfficeError("CONVERSION_UNSUPPORTED")


def merge_documents(
    sources: tuple[StructuredDocument, ...],
    target: DocumentFormat,
    limits: OfficeLimits,
) -> StructuredDocument:
    """Merge only explicitly selected one-sheet workbooks with identical typed headers."""
    if target is not DocumentFormat.XLSX or any(
        item.format is not DocumentFormat.XLSX
        or item.support is not DocumentSupport.EDITABLE
        or len(item.sheets) != 1
        or item.formula_count
        for item in sources
    ):
        raise OfficeError("MERGE_REQUIRES_SIMPLE_WORKBOOKS")
    first = sources[0].sheets[0]
    headers = tuple((cell.column, cell.value) for cell in first.cells if cell.row == 1)
    cells: list[DocumentCell] = list(first.cells)
    rows = first.rows
    for document in sources[1:]:
        sheet = document.sheets[0]
        if sheet.columns != first.columns or headers != tuple(
            (cell.column, cell.value) for cell in sheet.cells if cell.row == 1
        ):
            raise OfficeError("SCHEMA_MAPPING_REVIEW_REQUIRED")
        for cell in sheet.cells:
            if cell.row == 1:
                continue
            row = rows + cell.row - 1
            cells.append(
                cell.model_copy(
                    update={
                        "row": row,
                        "reference": f"s:0:{get_column_letter(cell.column)}{row}",
                    }
                )
            )
        rows += max(0, sheet.rows - 1)
    if rows > limits.max_rows or len(cells) > limits.max_cells:
        raise OfficeError("MERGE_LIMIT")
    return StructuredDocument(
        format=target,
        sheets=(
            first.model_copy(
                update={
                    "cells": tuple(cells),
                    "rows": rows,
                    "name": "Summary",
                }
            ),
        ),
    )


def apply_operation(
    document: StructuredDocument, operation: DocumentOperation
) -> StructuredDocument:
    """Dispatch one finite edit; unsupported combinations reject rather than guess."""
    kind = operation.kind
    text_kinds = {
        DocumentOperationKind.REPLACE_TEXT,
        DocumentOperationKind.APPEND_PARAGRAPH,
        DocumentOperationKind.SET_HEADING,
    }
    if kind in text_kinds:
        if document.format not in {
            DocumentFormat.TXT,
            DocumentFormat.MARKDOWN,
            DocumentFormat.DOCX,
        }:
            raise OfficeError("TEXT_OPERATION_FORMAT_MISMATCH")
        return edit_blocks(document, operation)
    if kind is DocumentOperationKind.SET_JSON_VALUE:
        if document.format is not DocumentFormat.JSON:
            raise OfficeError("JSON_OPERATION_FORMAT_MISMATCH")
        return document.model_copy(
            update={
                "blocks": (
                    DocumentBlock(
                        reference="body",
                        text=json_set(
                            document.blocks[0].text,
                            operation.target,
                            operation.value,
                        ),
                    ),
                )
            }
        )
    if kind in {DocumentOperationKind.ADD_WORKSHEET, DocumentOperationKind.RENAME_WORKSHEET}:
        if document.format is not DocumentFormat.XLSX:
            raise OfficeError("WORKSHEET_FORMAT_MISMATCH")
        return edit_sheet_names(document, operation)
    if kind in {
        DocumentOperationKind.FILTER_EMPTY_ROWS,
        DocumentOperationKind.SORT_ROWS,
        DocumentOperationKind.RENAME_COLUMN,
    }:
        if document.format is not DocumentFormat.CSV:
            raise OfficeError("CSV_OPERATION_FORMAT_MISMATCH")
        return edit_csv_rows(document, operation)
    return edit_cells(document, operation)


def edit_blocks(document: StructuredDocument, operation: DocumentOperation) -> StructuredDocument:
    """Apply exact whole-block changes; never guess a repeated heading or section."""
    if operation.value.kind is not ValueKind.TEXT:
        raise OfficeError("TEXT_VALUE_REQUIRED")
    blocks = list(document.blocks)
    if operation.kind is DocumentOperationKind.APPEND_PARAGRAPH:
        if operation.target != "end":
            raise OfficeError("APPEND_TARGET_MUST_BE_END")
        if document.format in {DocumentFormat.TXT, DocumentFormat.MARKDOWN}:
            blocks[0] = blocks[0].model_copy(
                update={
                    "text": blocks[0].text + document.newline + operation.value.value,
                }
            )
        else:
            blocks.append(
                DocumentBlock(
                    reference=f"p:{len(blocks)}",
                    text=operation.value.value,
                    heading_level=operation.heading_level,
                )
            )
    else:
        matches = [
            index for index, block in enumerate(blocks) if block.reference == operation.target
        ]
        if len(matches) != 1:
            raise OfficeError("DOCUMENT_TARGET_AMBIGUOUS_OR_MISSING")
        index = matches[0]
        if operation.expected is not None and operation.expected != OfficeValue(
            value=blocks[index].text
        ):
            raise OfficeError("EDIT_EXPECTED_VALUE_MISMATCH")
        update: dict[str, object] = {"text": operation.value.value}
        if operation.kind is DocumentOperationKind.SET_HEADING:
            if document.format is not DocumentFormat.DOCX or not operation.heading_level:
                raise OfficeError("HEADING_LEVEL_REQUIRED")
            update["heading_level"] = operation.heading_level
        blocks[index] = blocks[index].model_copy(update=update)
    return document.model_copy(update={"blocks": tuple(blocks)})


def edit_cells(document: StructuredDocument, operation: DocumentOperation) -> StructuredDocument:
    """Edit one unmerged unprotected exact cell and preserve formula ownership."""
    if document.format not in {DocumentFormat.DOCX, DocumentFormat.CSV, DocumentFormat.XLSX}:
        raise OfficeError("CELL_OPERATION_FORMAT_MISMATCH")
    if (document.format is DocumentFormat.DOCX) != (
        operation.kind is DocumentOperationKind.UPDATE_TABLE_CELL
    ):
        raise OfficeError("CELL_OPERATION_FORMAT_MISMATCH")
    if operation.kind not in {
        DocumentOperationKind.SET_SPREADSHEET_CELL,
        DocumentOperationKind.UPDATE_TABLE_CELL,
        DocumentOperationKind.SCALE_NUMBER,
    }:
        raise OfficeError("OPERATION_NOT_SUPPORTED")
    sheets = list(document.sheets)
    found = False
    for index, sheet in enumerate(sheets):
        cells = list(sheet.cells)
        for cell_index, cell in enumerate(cells):
            if cell.reference != operation.target:
                continue
            if found or sheet.protected:
                raise OfficeError("CELL_AMBIGUOUS_OR_PROTECTED")
            found = True
            for merged in sheet.merged_ranges:
                left, top, right, bottom = range_boundaries(merged)
                if left is None or top is None or right is None or bottom is None:
                    raise OfficeError("MERGED_CELL_EDIT_BLOCKED")
                if left <= cell.column <= right and top <= cell.row <= bottom:
                    raise OfficeError("MERGED_CELL_EDIT_BLOCKED")
            if operation.expected is not None and operation.expected != cell.value:
                raise OfficeError("EDIT_EXPECTED_VALUE_MISMATCH")
            value = operation.value
            if operation.kind is DocumentOperationKind.SCALE_NUMBER:
                if cell.value.kind is not ValueKind.NUMBER or value.kind is not ValueKind.NUMBER:
                    raise OfficeError("NUMERIC_CELL_REQUIRED")
                with localcontext() as context:
                    context.prec = 256
                    number = (Decimal(cell.value.value) * Decimal(value.value)).quantize(
                        Decimal("0.01"), rounding=ROUND_HALF_EVEN
                    )
                value = OfficeValue(kind=ValueKind.NUMBER, value=str(number))
            if cell.value.kind is ValueKind.FORMULA and value.kind is not ValueKind.FORMULA:
                raise OfficeError("FORMULA_CELL_REQUIRES_REVIEW")
            if value.kind is ValueKind.FORMULA:
                require_safe_formula(value.value)
            cells[cell_index] = cell.model_copy(update={"value": value})
        sheets[index] = sheet.model_copy(update={"cells": tuple(cells)})
    if not found:
        raise OfficeError("CELL_TARGET_MISSING")
    return document.model_copy(update={"sheets": tuple(sheets)})


def edit_sheet_names(
    document: StructuredDocument, operation: DocumentOperation
) -> StructuredDocument:
    """Reject name collisions and reference-sensitive rename rather than corrupt formulas."""
    name = operation.value.value
    if (
        operation.value.kind is not ValueKind.TEXT
        or not name
        or len(name) > 31
        or any(char in name for char in "[]:*?/\\")
        or name.startswith("'")
        or name.endswith("'")
        or any(sheet.name.casefold() == name.casefold() for sheet in document.sheets)
    ):
        raise OfficeError("INVALID_OR_CONFLICTING_SHEET_NAME")
    sheets = list(document.sheets)
    if operation.kind is DocumentOperationKind.ADD_WORKSHEET:
        sheets.append(DocumentSheet(reference=f"s:{len(sheets)}", name=name, rows=1, columns=1))
    else:
        if document.formula_count:
            raise OfficeError("SHEET_RENAME_REFERENCE_REVIEW_REQUIRED")
        matches = [
            index for index, sheet in enumerate(sheets) if sheet.reference == operation.target
        ]
        if len(matches) != 1:
            raise OfficeError("SHEET_TARGET_MISSING")
        sheets[matches[0]] = sheets[matches[0]].model_copy(update={"name": name})
    return document.model_copy(update={"sheets": tuple(sheets)})


def edit_csv_rows(document: StructuredDocument, operation: DocumentOperation) -> StructuredDocument:
    """Filter/sort explicitly selected CSV rows while keeping its header and all bad data."""
    sheet = document.sheets[0]
    indexed = {(cell.row, cell.column): cell.value for cell in sheet.cells}
    rows = [
        [indexed.get((row, column), OfficeValue()) for column in range(1, sheet.columns + 1)]
        for row in range(1, sheet.rows + 1)
    ]
    if not rows:
        raise OfficeError("CSV_EMPTY")
    if operation.kind is DocumentOperationKind.FILTER_EMPTY_ROWS:
        rows = [rows[0], *(row for row in rows[1:] if any(value.value.strip() for value in row))]
    else:
        if not operation.target.isdecimal() or not 1 <= int(operation.target) <= sheet.columns:
            raise OfficeError("CSV_COLUMN_REQUIRED")
        column = int(operation.target) - 1
        if operation.kind is DocumentOperationKind.SORT_ROWS:
            rows = [rows[0], *sorted(rows[1:], key=lambda row: row[column].value)]
        else:
            rows[0][column] = operation.value
    cells = tuple(
        DocumentCell(
            reference=f"s:0:r:{row_index}:c:{column}",
            row=row_index,
            column=column,
            value=value,
        )
        for row_index, row in enumerate(rows, 1)
        for column, value in enumerate(row, 1)
    )
    return document.model_copy(
        update={
            "sheets": (
                sheet.model_copy(
                    update={
                        "cells": cells,
                        "rows": len(rows),
                    }
                ),
            )
        }
    )


def document_diff(
    before: StructuredDocument | None, after: StructuredDocument
) -> tuple[DocumentDifference, ...]:
    """Compare complete source-addressed values, types and sheet names for local Preview."""

    def values(document: StructuredDocument | None) -> dict[str, tuple[str, str]]:
        if document is None:
            return {}
        result = {
            block.reference: (block.text, f"heading:{block.heading_level}")
            for block in document.blocks
        }
        for sheet in document.sheets:
            result[sheet.reference + ":name"] = (sheet.name, "sheet")
            result[sheet.reference + ":rows"] = (str(sheet.rows), "row_count")
            result.update(
                {cell.reference: (cell.value.value, cell.value.kind.value) for cell in sheet.cells}
            )
        return result

    old, new = values(before), values(after)
    return tuple(
        DocumentDifference(
            reference=reference,
            before=old.get(reference, ("", "absent"))[0],
            after=new.get(reference, ("", "absent"))[0],
            before_type=old.get(reference, ("", "absent"))[1],
            after_type=new.get(reference, ("", "absent"))[1],
        )
        for reference in sorted(old.keys() | new.keys())
        if old.get(reference) != new.get(reference)
    )
