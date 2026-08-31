"""Static XLSX adapter preserving scalar types and never calculating formulas."""

from __future__ import annotations

import io
from datetime import date, datetime, time, timedelta
from decimal import Decimal

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from pc_manager_agent.config.office import OfficeLimits
from pc_manager_agent.domain.office_documents import (
    DocumentCell,
    DocumentFormat,
    DocumentSheet,
    DocumentSupport,
    OfficeError,
    OfficeValue,
    StructuredDocument,
    ValueKind,
)
from pc_manager_agent.safety.office.content import require_safe_formula


def scalar(value: object, *, formula: bool = False) -> OfficeValue:
    """Convert supported native scalars without conflating text and formulas."""
    if value is None:
        return OfficeValue(kind=ValueKind.EMPTY)
    if isinstance(value, bool):
        return OfficeValue(kind=ValueKind.BOOLEAN, value="true" if value else "false")
    if isinstance(value, datetime):
        return OfficeValue(kind=ValueKind.DATETIME, value=value.isoformat())
    if isinstance(value, date):
        return OfficeValue(kind=ValueKind.DATE, value=value.isoformat())
    if isinstance(value, (int, float, Decimal)):
        return OfficeValue(kind=ValueKind.NUMBER, value=str(value))
    if isinstance(value, str):
        return OfficeValue(kind=ValueKind.FORMULA if formula else ValueKind.TEXT, value=value)
    if isinstance(value, (time, timedelta)):
        raise OfficeError("UNSUPPORTED_TIME_CELL")
    raise OfficeError("UNSUPPORTED_CELL_TYPE")


def parse_xlsx(
    data: bytes, format_: DocumentFormat, warnings: tuple[str, ...], limits: OfficeLimits
) -> StructuredDocument:
    """Read actual cells, formulas and protection; do not refresh linked resources."""
    workbook = load_workbook(
        io.BytesIO(data),
        read_only=False,
        data_only=False,
        keep_links=False,
        keep_vba=False,
    )
    flags = set(warnings)
    sheets: list[DocumentSheet] = []
    formulas = 0
    budget = 0
    try:
        if len(workbook.worksheets) > limits.max_sheets:
            raise OfficeError("SHEET_LIMIT")
        for index, sheet in enumerate(workbook.worksheets):
            if sheet.max_row > limits.max_rows:
                raise OfficeError("ROW_LIMIT")
            budget += sheet.max_row * sheet.max_column
            if budget > limits.max_cells:
                raise OfficeError("CELL_LIMIT")
            cells: list[DocumentCell] = []
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = scalar(cell.value, formula=cell.data_type == "f")
                    if cell.data_type == "e":
                        flags.add("SPREADSHEET_ERROR_CELL")
                    if value.kind is ValueKind.FORMULA:
                        formulas += 1
                        try:
                            require_safe_formula(value.value)
                        except OfficeError:
                            flags.add("UNSUPPORTED_FORMULA")
                    cells.append(
                        DocumentCell(
                            reference=f"s:{index}:{cell.coordinate}",
                            row=cell.row,
                            column=cell.column,
                            value=value,
                            number_format=cell.number_format,
                        )
                    )
            sheets.append(
                DocumentSheet(
                    reference=f"s:{index}",
                    name=sheet.title,
                    rows=sheet.max_row,
                    columns=sheet.max_column,
                    cells=tuple(cells),
                    merged_ranges=tuple(str(item) for item in sheet.merged_cells.ranges),
                    protected=bool(sheet.protection.sheet),
                )
            )
            if sheet.protection.sheet:
                flags.add("PROTECTED_DOCUMENT")
        if workbook.security.lockStructure or workbook.security.lockWindows:
            flags.add("PROTECTED_DOCUMENT")
    finally:
        workbook.close()
    if format_ is DocumentFormat.XLSM:
        flags.add("MACRO_ENABLED_DOCUMENT")
    return StructuredDocument(
        format=format_,
        sheets=tuple(sheets),
        warnings=tuple(sorted(flags)),
        support=DocumentSupport.READ_ONLY if flags else DocumentSupport.EDITABLE,
        formula_count=formulas,
    )


def set_cell(
    sheet: Worksheet, row: int, column: int, value: OfficeValue, number_format: str
) -> None:
    """Assign a typed value; text beginning '=' explicitly remains text."""
    cell = sheet.cell(row, column)
    if value.kind is ValueKind.EMPTY:
        cell.value = None
    elif value.kind is ValueKind.NUMBER:
        number = Decimal(value.value)
        cell.value = int(number) if number == number.to_integral() else float(number)
    elif value.kind is ValueKind.BOOLEAN:
        cell.value = value.value == "true"
    elif value.kind is ValueKind.DATE:
        cell.value = date.fromisoformat(value.value)
    elif value.kind is ValueKind.DATETIME:
        cell.value = datetime.fromisoformat(value.value)
    elif value.kind is ValueKind.FORMULA:
        require_safe_formula(value.value)
        cell.value = value.value
    else:
        cell.value = value.value
        cell.data_type = "s"
    cell.number_format = number_format


def serialize_xlsx(document: StructuredDocument, original: bytes | None) -> bytes:
    """Write supported cells to an in-memory package, never to an original path."""
    workbook = (
        load_workbook(io.BytesIO(original), keep_links=False, data_only=False)
        if original is not None
        else Workbook()
    )
    if original is None:
        workbook.remove(workbook.worksheets[0])
    try:
        for index, source in enumerate(document.sheets):
            target = (
                workbook.worksheets[index]
                if index < len(workbook.worksheets)
                else workbook.create_sheet()
            )
            target.title = source.name
            for cell in source.cells:
                set_cell(target, cell.row, cell.column, cell.value, cell.number_format)
        if not workbook.worksheets:
            raise OfficeError("WORKBOOK_NEEDS_SHEET")
        stream = io.BytesIO()
        workbook.save(stream)
        return stream.getvalue()
    finally:
        workbook.close()
